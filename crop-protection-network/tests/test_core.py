import copy
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook

from cpn.config import API_URL, SITES, Settings
from cpn.normalize import SourceError, ingredients, matching_sites, normalize
from cpn.sources import DumpLinks, epa_url, iter_api, iter_workbook
from cpn.store import Store
from cpn.sync import collect, publish, sync

TODAY = date(2026, 10, 15)


def row(index=1, **kwargs):
    return {"reg_num": f"12345-{index}", "product_name": f"Fixture {index:06}",
            "status_group": "Active", "status_dt": "2020-01-01T00:00:00Z",
            "sites": "Corn (seed treatment)", "ais": "Thiamethoxam (060109/153719-23-4) - (20%), Fludioxonil (071503/131341-86-1) - (2%),",
            **kwargs}


class NormalizationTests(unittest.TestCase):
    def test_every_exact_site_is_included(self):
        self.assertEqual(len(SITES), 14)
        self.assertEqual(len(set(SITES.values())), 8)
        for site in SITES:
            with self.subTest(site=site):
                self.assertEqual(normalize(row(sites=site), TODAY)["sites"], [site])

    def test_embedded_comma_in_edible_soybean(self):
        site = "Soybeans, Edible (Seed and Pod Vegetable) (seed treatment)"
        self.assertEqual(matching_sites("Other use, " + site + ", Corn (seed treatment),"), ["Corn (seed treatment)", site])

    def test_broad_or_suffix_site_is_not_a_match(self):
        for site in ["Soybeans", "Sweet Corn (seed treatment)", "Wheat (foliar treatment)", "Corn (seed treatment) extra"]:
            self.assertIsNone(normalize(row(sites=site), TODAY))

    def test_status_and_inclusive_date(self):
        self.assertIsNotNone(normalize(row(status_dt="2026-10-15T23:59:59Z"), TODAY))
        self.assertIsNotNone(normalize(row(status_dt="2000-01-01", status_group="Active Products"), TODAY))
        self.assertIsNone(normalize(row(status_dt="2026-10-16"), TODAY))
        self.assertIsNone(normalize(row(status_group="Inactive"), TODAY))
        with self.assertRaises(SourceError):
            normalize(row(status_dt="unknown"), TODAY)

    def test_ingredient_internal_comma_kept_and_all_components_parsed(self):
        value = "Complex chemical, isomer (123456/111-22-3) - (.25%), Thiamethoxam (060109/153719-23-4) - (20%),"
        ai, warnings = ingredients(value)
        self.assertEqual(len(ai), 2)
        self.assertEqual(ai[0]["name"], "Complex chemical, isomer")
        self.assertEqual(ai[0]["percent"], .25)
        self.assertEqual(ai[0]["purposes"], ["Unclassified"])
        self.assertEqual(ai[1]["family"], "Neonicotinoid")
        self.assertFalse(warnings)

    def test_unparsed_text_is_preserved(self):
        ai, warnings = ingredients("Unknown compound, strain Q")
        self.assertEqual(ai[0]["name"], "Unknown compound, strain Q")
        self.assertIsNone(ai[0]["percent"])
        self.assertTrue(warnings)

    def test_distributor_label_is_explicit_base_record(self):
        p = normalize(row(reg_num="12345-1-999"), TODAY)
        self.assertEqual(p["label_registration"], "12345-1")
        self.assertTrue(p["label_url"].endswith(":12345-1"))


class PaginationTests(unittest.TestCase):
    def test_more_than_ten_thousand_no_next_link_fallback(self):
        requested = []
        total = 10037
        def fetch(url):
            params = parse_qs(urlsplit(url).query)
            query = json.loads(params["q"][0])
            self.assertEqual(query["status_group"], "Active")
            offset = int(params.get("offset", [0])[0]); requested.append(offset)
            # Server clamps to 500 even when client requests 5000.
            items = [row(i) for i in range(offset, min(total, offset+500))]
            return {"items": items, "count": len(items), "hasMore": offset+len(items)<total,
                    "totalResults": total, "links": []}
        products, scanned = collect(iter_api(fetch), TODAY)
        self.assertEqual(len(products), total)
        self.assertEqual(scanned, total)
        self.assertEqual(requested[-1], 10000)

    def test_relative_next_link(self):
        def fetch(url):
            if "offset=1" in url:
                return {"items": [row(2)], "hasMore": False}
            return {"items": [row(1)], "hasMore": True, "links": [{"rel":"next", "href":"?offset=1"}]}
        self.assertEqual(len(list(iter_api(fetch))), 2)

    def test_broken_page_must_not_publish(self):
        for data in [
            {"items":[], "hasMore":True},
            {"items":[row()], "hasMore":False, "totalResults":2},
            {"items":[row()], "hasMore":False, "count":3},
            {"items":[row()]},
        ]:
            with self.subTest(data=data), self.assertRaises(SourceError):
                list(iter_api(lambda url: data))

    def test_duplicate_page_is_an_error(self):
        with self.assertRaises(SourceError):
            list(iter_api(lambda url: {"items":[row()], "hasMore":True}))

    def test_next_link_must_stay_on_epa_endpoint(self):
        with self.assertRaises(SourceError):
            list(iter_api(lambda url: {"items":[row()], "hasMore":True, "links":[{"rel":"next","href":"https://example.com/steal"}]}))
        with self.assertRaises(SourceError):
            epa_url("https://epa.gov.example.com/file.xlsx")


class WorkbookAndStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)
        self.store = Store(self.path / "test.sqlite3")

    def tearDown(self):
        self.temp.cleanup()

    def test_multisheet_workbook_all_rows_and_export_aliases(self):
        workbook = Workbook()
        for j in range(2):
            sheet = workbook.active if j==0 else workbook.create_sheet("Part 2")
            sheet.append(["Full APPRIL data dump"])
            sheet.append(["Registration Number", "Product Name", "Status Group", "Status Date", "Sites", "Active Ingredients"])
            for i in range(601):
                r=row(j*601+i)
                sheet.append([r[k] for k in ["reg_num","product_name","status_group","status_dt","sites","ais"]])
        file=self.path/"weekly.xlsx";workbook.save(file)
        products, scanned=collect(iter_workbook(file), TODAY)
        self.assertEqual(len(products),1202)
        self.assertEqual(scanned,1202)

    def test_incomplete_results_export_is_rejected(self):
        workbook=Workbook();workbook.active.append(["Registration Number","Product Name","Status Group"])
        workbook.active.append(["123-1","Test","Active"])
        file=self.path/"incomplete.xlsx";workbook.save(file)
        with self.assertRaises(SourceError):list(iter_workbook(file))

    def test_atomic_rollback_and_complete_replacement(self):
        p=normalize(row(),TODAY)
        self.store.publish([p],source="initial")
        with self.assertRaises(Exception):self.store.publish([p,p],source="broken")
        self.assertEqual(self.store.metadata()["source"],"initial")
        self.assertEqual(self.store.search({})["total"],1)
        self.store.publish([normalize(row(2),TODAY)],source="new")
        self.assertIsNone(self.store.get(p["reg_num"]))

    def test_exact_crops_partial_ingredient_and_sql_input(self):
        self.store.publish([normalize(row(),TODAY),normalize(row(2, sites="Cereal crops (seed treatment)"),TODAY)])
        self.assertEqual(self.store.search({"crop":"Corn"})["total"],1)
        self.assertEqual(self.store.search({"ingredient":"THIAM"})["total"],2)
        self.assertEqual(self.store.search({"q":"%' OR 1=1 --"})["total"],0)
        self.assertEqual(self.store.search({"q":"%"})["total"],2)  # Literal % is present in composition.

    def test_unchanged_data_preserves_version(self):
        p=normalize(row(),TODAY)
        publish(self.store,[p],"API",TODAY,1)
        first=self.store.metadata()["version"]
        publish(self.store,[p],"API",TODAY,1)
        self.assertEqual(first,self.store.metadata()["version"])

    @patch('cpn.sync.download_dump')
    @patch('cpn.sync.iter_api')
    def test_api_failure_does_not_roll_back_to_excel(self, api, dump):
        p=normalize(row(),TODAY)
        publish(self.store,[p],"EPA APPRIL REST API",TODAY,1)
        self.store.set_metadata(api_last_success="2026-10-15T00:00:00+00:00",dump_valid_sha256="hash")
        version=self.store.metadata()["version"]
        api.side_effect=SourceError("source unavailable")
        dump.return_value=(self.path/"unused.xlsx",{"sha256":"hash"},False)
        result=sync(self.store,Settings(data_dir=self.path,demo=False))
        self.assertFalse(result["api_ok"])
        self.assertEqual(version,self.store.metadata()["version"])
        self.assertFalse(self.store.metadata()["sync_running"])

    def test_xlsx_discovery_parser(self):
        parser=DumpLinks();parser.feed('<a href="weekly.xlsx">Full <b>data dump</b></a>')
        self.assertEqual(parser.links,[["weekly.xlsx","Full data dump"]])


if __name__ == '__main__':unittest.main()
