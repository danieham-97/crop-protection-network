"""EPA adapters. No total-row cap; only per-page/transport safety bounds."""
import hashlib
import json
import time
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

from .config import API_URL, LANDING_URL
from .normalize import SourceError, canonical, text, validate_schema


def epa_url(url):
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme != "https" or not (parsed.hostname == "epa.gov" or (parsed.hostname or "").endswith(".epa.gov")) or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise SourceError("Source URLs must be HTTPS addresses on epa.gov")
    return url


class EPARedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return super().redirect_request(req, fp, code, msg, headers, epa_url(newurl))


def request(url, headers=None):
    opener = urllib.request.build_opener(EPARedirect)
    req = urllib.request.Request(epa_url(url), headers={
        "User-Agent": "CPN-Seed-Treatment-Explorer/1.0 (public EPA data cache)",
        "Accept": "application/json, application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, text/html",
        **(headers or {}),
    })
    for attempt in range(4):
        try:
            return opener.open(req, timeout=90)
        except urllib.error.HTTPError as exc:
            if exc.code == 304:
                return exc
            if exc.code not in (429, 500, 502, 503, 504) or attempt == 3:
                raise SourceError(f"EPA returned HTTP {exc.code}") from exc
            delay = exc.headers.get("Retry-After", "")
            time.sleep(min(30, int(delay)) if delay.isdigit() else 2 ** attempt)
        except (urllib.error.URLError, TimeoutError) as exc:
            if attempt == 3:
                raise SourceError("EPA connection failed or timed out") from exc
            time.sleep(2 ** attempt)


def get_json(url):
    with request(url) as response:
        try:
            return json.load(response)
        except (ValueError, UnicodeError) as exc:
            raise SourceError("EPA did not return JSON") from exc


def iter_api(fetch=get_json, page_size=5000):
    query = {"status_group": "Active", "$orderby": {"ri_and_dp_num_sort": "asc"}}
    url = API_URL + "?" + urllib.parse.urlencode({"q": json.dumps(query, separators=(",", ":")), "limit": page_size})
    visited, registrations = set(), set()
    expected_total = None
    count = 0
    while url:
        epa_url(url)
        if url in visited:
            raise SourceError("EPA pagination repeated a page; snapshot was not published")
        visited.add(url)
        data = fetch(url)
        if not isinstance(data, dict) or not isinstance(data.get("items"), list):
            raise SourceError("Invalid ORDS response: items array is required")
        if not isinstance(data.get("hasMore"), bool):
            raise SourceError("Missing ORDS hasMore completeness indicator")
        rows = data["items"]
        if data.get("count") is not None and int(data["count"]) != len(rows):
            raise SourceError("EPA page count does not match its rows")
        if data.get("totalResults") is not None:
            if expected_total is not None and expected_total != int(data["totalResults"]):
                raise SourceError("EPA total changed during pagination; retry required")
            expected_total = int(data["totalResults"])
        for row in rows:
            validate_schema(row)
            reg = text(canonical(row)["reg_num"])
            if reg in registrations:
                raise SourceError("Duplicate registration across EPA pages; retry required")
            registrations.add(reg)
            count += 1
            yield row
        next_links = [link.get("href") for link in data.get("links", []) if link.get("rel") == "next"]
        if data["hasMore"]:
            if not rows:
                raise SourceError("EPA returned an empty page with hasMore=true")
            if next_links:
                next_url = urllib.parse.urljoin(url, next_links[0])
                if urllib.parse.urlsplit(next_url).path != urllib.parse.urlsplit(API_URL).path:
                    raise SourceError("EPA next link changed the API endpoint")
                url = next_url
            else:
                # ORDS offset is based on actual rows, not the requested page size.
                url = API_URL + "?" + urllib.parse.urlencode({"q": json.dumps(query, separators=(",", ":")), "limit": page_size, "offset": count})
        else:
            if next_links:
                raise SourceError("Conflicting EPA pagination indicators")
            url = None
    if count == 0:
        raise SourceError("EPA active-products collection unexpectedly returned no records")
    if expected_total is not None and count != expected_total:
        raise SourceError("EPA collection ended before the advertised total")


class DumpLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links, self.current = [], None

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.current = [dict(attrs).get("href", ""), ""]

    def handle_data(self, data):
        if self.current is not None:
            self.current[1] += data

    def handle_endtag(self, tag):
        if tag == "a" and self.current is not None:
            self.links.append(self.current)
            self.current = None


def discover_dump():
    with request(LANDING_URL) as response:
        parser = DumpLinks()
        parser.feed(response.read().decode("utf-8", errors="replace"))
    candidates = [(href, label) for href, label in parser.links
                  if "data dump" in label.lower() or "data dump" in href.lower() or ".xlsx" in href.lower()]
    candidates.sort(key=lambda link: "data dump" not in link[1].lower())
    for href, _ in candidates:
        if href and not href.lower().startswith("javascript:"):
            return epa_url(urllib.parse.urljoin(LANDING_URL, href))
    raise SourceError("Weekly Excel download link not found; set EPA_DUMP_URL to the full data dump link on APPRIL")


def download_dump(data_dir, override=""):
    data_dir = Path(data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)
    destination = data_dir / "appril-weekly.xlsx"
    metadata_file = data_dir / "appril-weekly.json"
    try:
        previous = json.loads(metadata_file.read_text()) if metadata_file.exists() else {}
    except (ValueError, OSError):
        previous = {}  # Recover a metadata write interrupted by host shutdown.
    url = override or discover_dump()
    headers = {}
    if destination.exists() and previous.get("url") == url:
        if previous.get("etag"):
            headers["If-None-Match"] = previous["etag"]
        if previous.get("last_modified"):
            headers["If-Modified-Since"] = previous["last_modified"]
    temporary = destination.with_suffix(".part")
    try:
        with request(url, headers) as response:
            if response.code == 304:
                if not destination.exists():
                    raise SourceError("EPA returned 304 without a local workbook")
                return destination, previous, False
            digest = hashlib.sha256()
            with temporary.open("wb") as output:
                while chunk := response.read(1024 * 1024):
                    output.write(chunk)
                    digest.update(chunk)
            with temporary.open("rb") as f:
                if f.read(2) != b"PK":
                    raise SourceError("Weekly download is not an XLSX workbook")
            metadata = {"url": url, "etag": response.headers.get("ETag"),
                        "last_modified": response.headers.get("Last-Modified"), "sha256": digest.hexdigest()}
        changed = previous.get("sha256") != metadata["sha256"]
        temporary.replace(destination)
        metadata_temporary = metadata_file.with_suffix(".tmp")
        metadata_temporary.write_text(json.dumps(metadata))
        metadata_temporary.replace(metadata_file)
        return destination, metadata, changed
    finally:
        temporary.unlink(missing_ok=True)


def iter_workbook(path):
    """Stream every row on every flat APPRIL data sheet, including split exports.

    A results-page export lacks Sites and is rejected. Relationship-sheet schemas
    are rejected explicitly instead of silently producing a partial catalog.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    found = False
    count = 0
    try:
        for sheet in workbook.worksheets:
            headers = None
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if headers is None:
                    candidate = canonical({str(v): None for v in values if v is not None})
                    if {"reg_num", "product_name"} <= set(candidate):
                        headers = [text(v) for v in values]
                        validate_schema(dict.fromkeys(headers))
                        found = True
                    elif number >= 30:
                        break  # Only header discovery is bounded; data rows are unlimited.
                    continue
                if not any(v is not None for v in values):
                    continue
                row = dict(zip(headers, values))
                if not text(canonical(row).get("reg_num")):
                    raise SourceError(f"Missing registration in {sheet.title}, row {number}")
                count += 1
                yield row
        if not found or count == 0:
            raise SourceError("No complete APPRIL data sheets found in weekly workbook")
    finally:
        workbook.close()
